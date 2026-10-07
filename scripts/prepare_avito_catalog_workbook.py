"""Build review-only Avito XLSX and balanced channel copy from sanitized live evidence."""

import argparse
from collections import Counter
from copy import copy
from datetime import date
import hashlib
import json
from pathlib import Path
import re
from urllib.parse import urlsplit

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils.cell import range_boundaries


TEMPLATE_SHA256 = "8f2c75e5dab0801a0b3b4693e58df82f73f09563db9a837a479e68b338d7c30c"
PHONE_SHEET = "Телефоны-Мобильные телефоны"
COLORS = {"Deep Purple": "фиолетовый", "Silver": "серебристый", "Space Black": "черный",
          "Black Titanium": "черный", "White Titanium": "белый", "Blue Titanium": "синий",
          "Natural Titanium": "серый", "Burgundy": "красный",
          "Cream": "бежевый", "Titanium Black": "черный", "Gold": "золотистый"}
# Owner's validator rejected beige; owner then confirmed gold in the model's
# Color dictionary. Full XML/account validation remains a separate gate.
MODEL_COLORS = {("iPhone 16 Pro Max", "Desert Titanium"): "золотистый"}


def avito_color(model, color):
    return MODEL_COLORS.get((model, color), COLORS.get(color))


LABELS = {
    "not_published_ready": "Черновик/архив или карточка не ready",
    "unsupported_phone_category": "Не смартфон: требуется отдельный шаблон Avito",
    "unsupported_device_condition": "Для новой техники нужен отдельный формат описания",
    "sold_or_unavailable": "Продано или нет доступного остатка",
    "inventory_link_count": "Нет однозначной связи со складом",
    "inventory_sku_mismatch": "SKU карточки не совпадает со складской строкой",
    "inventory_identity_unmatched": "Нет подтверждённого сопоставления с поступлением",
    "inventory_not_approved": "Нет полного допуска складской позиции",
    "inventory_sale_duplicate_or_issue": "Запрет продажи, дубль serial или открытая проблема",
    "inventory_stock_shortfall": "На складе меньше, чем на сайте",
    "inventory_retail_price_differs_from_site": "Цена snapshot отличается от цены сайта",
    "sold_but_inventory_snapshot_positive": "Продано на сайте, но в старом snapshot остаток положительный",
    "serial_present_identity_not_applicable_review": "Есть serial, но identity=not_applicable: проверить основание",
    "store_offer_mismatch": "Предложение магазина не совпадает с карточкой",
    "missing_details_passport_or_public_report": "Нет полного Passport/деталей/публичного сертификата",
    "diagnostic_date_mismatch": "Дата диагностики отличается от даты сертификата",
    "unknown_battery": "Неизвестно состояние батареи",
    "diagnostic_is_historical_recheck_before_publication": "Августовская диагностика: подтвердить актуальность",
    "public_page_photo_or_certificate_failed": "Не прошла проверка страницы, фото или сертификата",
    "gallery_has_five_not_six_images": "Опубликовано пять фото вместо шести: не блокирует экспорт",
    "duplicate_product_or_sku": "Дубль товара или SKU",
    "photo_asset_shared_with_other_product": "Фото-файл используется у другого товара",
    "channel_listing_missing_or_duplicate": "Нет однозначного канального объявления",
    "duplicate_external_id": "Дубль ID Avito",
    "description_missing": "Недостаточно данных для описания",
    "official_catalog_fields_and_xml_qa_pending": "Не завершена проверка справочника/обязательных полей/XML Avito",
    "channel_mapping_not_confirmed": "Соответствие категории Avito не подтверждено",
    "invalid_channel_price": "Некорректная цена канала",
}


def labels(codes):
    return "; ".join(LABELS.get(code, code) for code in codes)


def one(values):
    return values[0] if isinstance(values, list) and len(values) == 1 else {}


def percentage(value):
    match = re.fullmatch(r"(\d{1,3})%?", str(value if value is not None else "").strip())
    return int(match[1]) if match and 0 <= int(match[1]) <= 100 else None


def cycle_text(value):
    ending = "цикл" if value % 10 == 1 and value % 100 != 11 else (
        "цикла" if value % 10 in (2, 3, 4) and value % 100 not in (12, 13, 14) else "циклов")
    return f"{value} {ending}"


def image_urls(product):
    images = sorted(product.get("images") or [], key=lambda i: (i["role"] != "card", i["sort"] or 0))
    ids = list(dict.fromkeys([product.get("listing_file")] + [i["id"] for i in images]))
    valid = [i for i in ids if isinstance(i, str) and re.fullmatch(
        r"[0-9a-f]{8}(?:-[0-9a-f]{4}){3}-[0-9a-f]{12}", i, re.I)]
    return ["https://api.isvoi.ru/assets/" + i + "?width=1600&height=1600&fit=inside&format=jpg&quality=90"
            for i in valid[:10]]


def assess(product, public):
    blockers, warnings = [], []
    if product["status"] != "published" or product["content_status"] != "ready":
        blockers.append("not_published_ready")
    if product.get("category", {}).get("slug") != "smartphones":
        blockers.append("unsupported_phone_category")
    if product.get("condition") != "used":
        blockers.append("unsupported_device_condition")
    if product["stock_status"] != "available" or product["stock_quantity"] <= 0:
        blockers.append("sold_or_unavailable")
    inventory = one(product.get("inventory"))
    if not inventory:
        blockers.append("inventory_link_count")
    else:
        if inventory["source_sku"] != product["sku"]:
            blockers.append("inventory_sku_mismatch")
        if inventory["identity_status"] not in ("matched", "not_applicable"):
            blockers.append("inventory_identity_" + inventory["identity_status"])
        if inventory["eligibility_status"] != "eligible" or inventory["authenticity_status"] not in (
                "verified", "not_required") or not inventory["review_override"] or not inventory["review_note_present"]:
            blockers.append("inventory_not_approved")
        if not inventory["for_sale"] or inventory["duplicate_serial"] or inventory["open_issue_codes"]:
            blockers.append("inventory_sale_duplicate_or_issue")
        if inventory["quantity"] < product["stock_quantity"]:
            blockers.append("inventory_stock_shortfall")
        if inventory["retail_price"] != product["price"]:
            warnings.append("inventory_retail_price_differs_from_site")
        if product["stock_status"] == "sold" and inventory["quantity"] > 0:
            warnings.append("sold_but_inventory_snapshot_positive")
        if inventory["identity_status"] == "not_applicable" and inventory["serial_present"]:
            warnings.append("serial_present_identity_not_applicable_review")
    offers = [o for o in product.get("offers") or [] if o["status"] == "published" and o["location"] == "belgorod"]
    offer = one(offers)
    if not offer or any(offer.get(k) != product[k] for k in ("price", "stock_status", "stock_quantity")):
        blockers.append("store_offer_mismatch")
    details = one(product.get("device_details"))
    passport = one(product.get("passports"))
    report = one([r for r in product.get("reports") or [] if r["status"] == "current"])
    if not details or not passport or not report.get("public_file"):
        blockers.append("missing_details_passport_or_public_report")
    else:
        if str(report["tested_at"])[:10] != details.get("diagnostic_date"):
            blockers.append("diagnostic_date_mismatch")
        if percentage(details.get("battery")) is None:
            blockers.append("unknown_battery")
        warnings.append("diagnostic_is_historical_recheck_before_publication")
    if product["status"] == "published":
        page = next((p for p in public.get("pages", []) if p["product_id"] == product["id"]), {})
        photos = next((p for p in public.get("images", {}).get("products", []) if p["product_id"] == product["id"]), {})
        certs = [c for c in public.get("certificates", []) if c["sku"] == product["sku"]]
        if not page.get("ok") or not photos.get("ok") or len(certs) != 1 or not certs[0]["ok"]:
            blockers.append("public_page_photo_or_certificate_failed")
    if len(product.get("images") or []) < 6 and product["status"] == "published":
        warnings.append("gallery_has_five_not_six_images")
    return blockers, warnings


def balanced_copy(product, address):
    if product.get("condition") != "used":
        return "", ""
    d = one(product.get("device_details"))
    p = one(product.get("passports"))
    r = one([r for r in product.get("reports") or [] if r["status"] == "current"])
    if (not d or not p or not r or not product.get("completeness") or not product.get("warranty_text")
            or percentage(d.get("battery")) is None or not d.get("diagnostic_date")):
        return "", ""
    title = " ".join(str(x) for x in [product["model"], d["storage"], product["color"]] if x)
    if "Galaxy" in title and not title.startswith("Samsung"):
        title = "Samsung " + title
    condition = p.get("condition_note") or "Состояние показано на реальных фотографиях."
    notes = p.get("condition_notes") or []
    explicit_body = next((n for n in notes if re.search(r"нет замечаний по корпусу|замечаний по корпусу.*не обнаружено", n, re.I)), "")
    if condition.startswith("Диагностика пройдена, состояние корпуса показано"):
        condition = explicit_body.rstrip(".") + "." if explicit_body else "Состояние корпуса показано на фотографиях."
    if explicit_body and "корпус" not in condition.lower():
        condition = explicit_body.rstrip(".") + ". " + condition
    grade = "" if "грейд" in condition.lower() else f"Грейд {d['grade']}. "
    display_model = "Samsung " + product["model"] if "Galaxy" in product["model"] and not product["model"].startswith("Samsung") else product["model"]
    intro = f"{display_model}, {d['storage']}, {product['color']}. Устройство с пробегом. {grade}{condition} На фото именно этот телефон."
    when = date.fromisoformat(d["diagnostic_date"]).strftime("%d.%m.%Y")
    cycles = f", {cycle_text(d['battery_cycles'])}" if isinstance(d.get("battery_cycles"), int) else ""
    battery = f"Батарея: {percentage(d['battery'])}%{cycles} по отчёту {r['provider']} от {when}."
    checklist = p.get("diagnostics_checklist") or []
    checks = [c["text"] for c in checklist if c.get("state") == "ok" and not re.search(
        r"гарнитур|блокиров|Find My|MDM|FRP|ROOT|KNOX|Erase|Оригиналь", c.get("text", ""), re.I)]
    verification = "По диагностике прошли проверку: " + ", ".join(c[:1].lower() + c[1:] for c in checks) + "." if checks else ""
    manual_headset = any("гарнитур" in c.get("text", "").lower() for c in checklist)
    manual_note = ("Гарнитура не подключалась к стенду NSYS; её работоспособность подтверждена отдельно."
                   if manual_headset and ("не подключ" in condition.lower() or "не подключ" in (r.get("public_note") or "").lower()) else "")
    remarks = [c["text"] for c in checklist if c.get("state") not in ("ok", None)]
    if remarks:
        verification += " Проверки с замечаниями: " + ", ".join(remarks) + "."
    repair = p.get("repair") or ""
    if repair == "Проверенные компоненты отмечены диагностикой как оригинальные.":
        repair = "Проверенные компоненты отмечены как оригинальные."
    story = ". ".join(p.get("story_facts") or [])
    if story:
        story += "."
    warranty = product["warranty_text"]
    if re.fullmatch(r"Гарантия 90 дней\.?", warranty):
        warranty = "Гарантия магазина 90 дней"
    contents = f"Комплект: {product['completeness'].lower().rstrip('.')}. {warranty.rstrip('.')}."
    passport = "Есть Passport с результатами проверки. Публичная выписка скрывает полные идентификаторы; полный сертификат доступен в магазине."
    if address == 'Белгород, ул. Костюкова д.67А ТЦ "Виктория", 2 этаж':
        address = "Белгород, ул. Костюкова, д. 67А, ТЦ «Виктория», 2 этаж"
    cta = "Посмотреть устройство можно в I СВОИ: " + address.rstrip(".") + ". Напишите в чат: уточним наличие и согласуем время просмотра."
    description = "\n\n".join(x for x in [intro, battery, " ".join(x for x in [verification, repair, story] if x),
                                            manual_note, contents, passport, cta] if x)
    if len(title) > 50 or len(description) > 7500 or re.search(r"[<>]|\d{15}|\[redacted", title + description):
        raise ValueError("Unsafe or oversized channel copy")
    return title, description


def extra_sheet(workbook, name, headers, rows, widths):
    sheet = workbook.create_sheet(name)
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for row in sheet:
        for cell in row:
            cell.font = Font(name="Arial", size=10, bold=cell.row == 1,
                             color="FFFFFF" if cell.row == 1 else "222222")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if cell.row == 1:
                cell.fill = PatternFill("solid", fgColor="254A68")
            if cell.data_type == "f":
                cell.data_type = "s"
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = width
    return sheet


def prepare(snapshot, public, approved):
    products = snapshot["products"]
    stores = [s for s in snapshot["stores"] if s["slug"] == "belgorod" and s["status"] == "published"]
    store = one(stores)
    if not store.get("address") or snapshot.get("private_identifier_redactions"):
        raise ValueError("Missing store or redacted public source; operator review required")
    by_sku = Counter(p["sku"] for p in products)
    by_id = Counter(p["id"] for p in products)
    external_ids = Counter(l["external_id"] for p in products for l in p.get("listings") or [])
    assets = {}
    for p in products:
        for url in image_urls(p):
            assets.setdefault(url, set()).add(p["id"])
    shared_photos = {url for url, owners in assets.items() if len(owners) > 1}
    audits, entries = [], []
    for p in products:
        blockers, warnings = assess(p, public)
        if by_sku[p["sku"]] > 1 or by_id[p["id"]] > 1:
            blockers.append("duplicate_product_or_sku")
        if any(url in shared_photos for url in image_urls(p)):
            warnings.append("photo_asset_shared_with_other_product")
        listings = p.get("listings") or []
        listing = one(listings)
        if len(listings) != 1 or not listing.get("external_id"):
            blockers.append("channel_listing_missing_or_duplicate")
        elif external_ids[listing["external_id"]] > 1:
            blockers.append("duplicate_external_id")
        effective_price = listing.get("price_override") if listing.get("price_override") is not None else p["price"]
        if (isinstance(effective_price, bool) or not isinstance(effective_price, (int, float))
                or effective_price <= 0 or effective_price != int(effective_price)):
            blockers.append("invalid_channel_price")
        audit = {"sku": p["sku"], "product_id": p["id"], "title": p["title"], "status": p["status"],
                 "stock_status": p["stock_status"], "stock_quantity": p["stock_quantity"], "price": p["price"],
                 "inventory_identity_status": one(p.get("inventory")).get("identity_status"),
                 "inventory_quantity": one(p.get("inventory")).get("quantity"),
                 "photos": len(image_urls(p)), "blockers": blockers, "warnings": warnings,
                 "catalog_checks_pass": not blockers}
        audits.append(audit)
        if p["status"] != "published" or p["stock_status"] != "available" or p["stock_quantity"] <= 0:
            continue
        title, description = balanced_copy(p, store["address"])
        if not description:
            audit["blockers"].append("description_missing")
            audit["catalog_checks_pass"] = False
        approval = False
        previous = next((e for e in approved.get("entries", []) if e["sku"] == p["sku"]), {})
        if previous and approved.get("copy_operator_approved"):
            # Approval of a previous text is not a license to silently change its device facts.
            d = one(p.get("device_details"))
            date_text = date.fromisoformat(d["diagnostic_date"]).strftime("%d.%m.%Y")
            old = previous["description_override"]
            condition_source = " ".join([one(p.get("passports")).get("condition_note") or ""]
                                        + (one(p.get("passports")).get("condition_notes") or []))
            same_approved_condition = (
                f"грейд {d['grade']}".lower() in old.lower()
                and not re.search(r"царап|трещ|скол|потёр|потер|щербин|следы использования", condition_source, re.I)
                and one(p.get("passports")).get("repair") == "Проверенные компоненты отмечены диагностикой как оригинальные."
                and all(c.get("state") == "ok" for c in one(p.get("passports")).get("diagnostics_checklist") or []))
            if (same_approved_condition and previous.get("external_id") == listing.get("external_id")
                    and old.startswith(f"{p['model']}, {d['storage']}, {p['color']}.")
                    and date_text in old
                    and re.search(rf"(?<!\d){percentage(d['battery'])}%", old)
                    and re.search(rf"(?<!\d){d['battery_cycles']} цикл", old)):
                description, approval = old, True
        entries.append({"sku": p["sku"], "product_id": p["id"], "channel": "avito",
                        "external_id": listing.get("external_id"), "title_override": title,
                        "description_override": description, "price_rub": effective_price,
                        "price_source": "channel_override" if listing.get("price_override") is not None else "site",
                        "stock_quantity": p["stock_quantity"], "image_urls": image_urls(p),
                        "identity_status": audit["inventory_identity_status"],
                        "catalog_checks_pass": audit["catalog_checks_pass"],
                        "copy_operator_approved": approval, "ready_to_publish": False,
                        "blockers": audit["blockers"] + ["official_catalog_fields_and_xml_qa_pending",
                                                          "channel_mapping_not_confirmed"],
                        "warnings": warnings})
    return {"captured_at": snapshot["captured_at"], "selected_variant": "balanced", "apply": False,
            "ready_to_publish": False, "publication_performed": False, "pilot_limit_unchanged": 3,
            "products_audited": len(products), "published": sum(p["status"] == "published" for p in products),
            "available": len(entries), "catalog_candidate_count": sum(e["catalog_checks_pass"] for e in entries),
            "audits": audits, "entries": entries}


def apply_reference_evidence(report, snapshot, model_reference=None, box_confirmation=None):
    if model_reference:
        if (model_reference.get("scope") != "model_hardware_reference_not_avito_catalog_validation"
                or type(model_reference.get("version")) is not int or model_reference["version"] != 1
                or date.fromisoformat(model_reference["checked_on"]) > date.fromisoformat(snapshot["captured_at"][:10])):
            raise ValueError("Invalid model reference scope/version/date")
        for model, reference in model_reference.get("models", {}).items():
            attrs = reference.get("attributes") or {}
            source = urlsplit(reference.get("source_url") or "")
            if (not model or set(attrs) != {"RamSize"} or not re.fullmatch(r"[1-9][0-9]? ГБ", str(attrs["RamSize"]))
                    or source.scheme != "https" or source.hostname != "www.ifixit.com"
                    or not source.path.startswith("/Teardown/") or source.query or source.fragment
                    or source.username or source.password or source.port not in (None, 443)):
                raise ValueError("Invalid model hardware reference")
    if box_confirmation and (
            box_confirmation.get("source") != "owner_chat_confirmation"
            or box_confirmation.get("confirmed_on") != snapshot["captured_at"][:10]
            or box_confirmation.get("scope") != "current_site_devices_used_and_open_boxes_only"
            or box_confirmation.get("devices_condition") != "used"
            or box_confirmation.get("boxes_sealed") is not False):
        raise ValueError("Invalid owner box confirmation")
    for entry in report["entries"]:
        product = next(p for p in snapshot["products"] if p["id"] == entry["product_id"])
        reference = (model_reference or {}).get("models", {}).get(product["model"])
        if reference:
            entry["model_reference_attributes"] = dict(reference["attributes"])
            entry["model_reference_source"] = reference["source_url"]
        if box_confirmation:
            if product["condition"] != "used":
                raise ValueError("Owner used-device statement conflicts with catalog condition")
            if "коробк" in (product.get("completeness") or "").lower():
                entry.setdefault("confirmed_phone_attributes", {})["BoxSealed"] = "Нет"
    report["model_reference"] = model_reference
    report["owner_box_confirmation"] = box_confirmation
    return report


def validate_template_choices(workbook):
    sheet = workbook[PHONE_SHEET]
    checks = 0
    for rule in sheet.data_validations.dataValidation:
        if rule.type != "list" or not rule.formula1 or "!" not in rule.formula1:
            continue
        name, area = rule.formula1.rsplit("!", 1)
        dictionary = workbook[name.strip("'")]
        first_col, first_row, last_col, last_row = range_boundaries(area)
        values = {dictionary.cell(row, col).value for row in range(first_row, last_row + 1)
                  for col in range(first_col, last_col + 1)}
        for row in range(5, sheet.max_row + 1):
            for col in range(1, 35):
                cell = sheet.cell(row, col)
                if cell.value is None or cell.coordinate not in rule.sqref:
                    continue
                # Y2 in the owner-supplied template explicitly documents this
                # delimiter for a multi-valued kit, despite the Excel list rule.
                choices = str(cell.value).split(" | ") if col == 25 else [cell.value]
                if len(set(choices)) != len(choices) or any(choice not in values for choice in choices):
                    raise ValueError("Value absent from official flat dictionary: " + cell.coordinate)
                checks += 1
    return checks


def write_workbook(template, target, report, snapshot):
    if hashlib.sha256(template.read_bytes()).hexdigest() != TEMPLATE_SHA256:
        raise ValueError("Unexpected template version")
    workbook = load_workbook(template)
    sheet = workbook[PHONE_SHEET]
    dictionary = workbook["Спр-Телефоны-Мобильные телефоны"]
    ram_values = {dictionary.cell(row, 23).value for row in range(2, 78)}
    styles = [copy(sheet.cell(5, column)._style) for column in range(1, 35)]
    sheet.delete_rows(5, sheet.max_row - 4)
    products = {p["id"]: p for p in snapshot["products"]}
    store = next(s for s in snapshot["stores"] if s["slug"] == "belgorod")
    for row, entry in enumerate([e for e in report["entries"] if e["catalog_checks_pass"]], 5):
        p = products[entry["product_id"]]
        d = one(p["device_details"])
        vendor = "Apple" if p["model"].startswith("iPhone") else "Samsung" if "Galaxy" in p["model"] else ""
        values = {1: entry["external_id"], 6: store["address"], 7: entry["title_override"],
                  8: entry["description_override"], 9: " | ".join(entry["image_urls"]),
                  12: "Телефоны", 13: entry["price_rub"], 14: "Мобильные телефоны",
                  15: "Товар приобретен на продажу", 16: "Б/у", 18: vendor,
                  19: p["model"], 20: d["storage"], 22: avito_color(p["model"], p["color"]) or "",
                  27: percentage(d["battery"]) if vendor == "Apple" else None, 28: "Включается"}
        values[23] = entry.get("model_reference_attributes", {}).get("RamSize")
        values[26] = entry.get("confirmed_phone_attributes", {}).get("BoxSealed")
        if values[23] and values[23] not in ram_values:
            raise ValueError("Model RAM value is absent from official template dictionary")
        contents = p["completeness"].lower()
        if "коробк" in contents and "кабел" in contents:
            values[25] = "Коробка | Провод зарядки"
        for column in range(1, 35):
            cell = sheet.cell(row, column)
            cell._style = copy(styles[column - 1])
            cell.value = values.get(column)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if cell.data_type == "f":
                raise ValueError("Formula injection in template output")
        for column in (18, 19, 20, 22, 23, 25, 26, 29, 30):
            cell = sheet.cell(row, column)
            cell.fill = PatternFill("solid", fgColor="FFF1C2")
            cell.comment = Comment("Требуется подтверждение оператором по текущему справочнику Avito и фактическому состоянию. Это файл подготовки, не готовая загрузка.", "ISVOI")
        if values[23]:
            sheet.cell(row, 23).comment = Comment(
                "Источник: " + entry["model_reference_source"] + "; сверено " + report["model_reference"]["checked_on"]
                + ". Характеристика модели, не измерение экземпляра. Подтвердить сочетание в каталоге Avito.", "ISVOI")
        if values[26]:
            sheet.cell(row, 26).fill = PatternFill("solid", fgColor="E2F0D9")
            sheet.cell(row, 26).comment = Comment(
                "Источник: подтверждение владельца в чате, " + report["owner_box_confirmation"]["confirmed_on"]
                + ". Все текущие устройства с пробегом, коробки открытые. XML-комплект ещё требует QA.", "ISVOI")
        sheet.cell(row, 8).comment = Comment("Источник: production Passport, " + report["captured_at"]
                                             + "; SKU " + entry["sku"] + ". Батарея указана с датой отчёта.", "ISVOI")
        sheet.cell(row, 13).number_format = '#,##0" ₽"'
        sheet.row_dimensions[row].height = 115
    sheet.freeze_panes = "G5"
    audit_rows = [[a["sku"], a["title"], a["status"], a["stock_status"], a["stock_quantity"],
                   a["inventory_quantity"], a["inventory_identity_status"], a["price"], a["photos"],
                   labels(a["blockers"]), labels(a["warnings"])] for a in report["audits"]]
    extra_sheet(workbook, "Сверка всех устройств",
                ["SKU", "Устройство", "Статус карточки", "Наличие", "Сайт, шт.", "Snapshot склада, шт.",
                 "Identity", "Цена, ₽", "Фото для Avito", "Блокеры каталога", "Предупреждения"],
                audit_rows, [12, 55, 20, 18, 14, 23, 20, 16, 18, 55, 65])
    copy_sheet = extra_sheet(workbook, "Все доступные описания",
                ["SKU", "Заголовок", "Описание balanced", "Проверки каталога", "Текст одобрен", "Блокеры запуска"],
                [[e["sku"], e["title_override"], e["description_override"],
                  "Пройдены" if e["catalog_checks_pass"] else "Блокировано",
                  "Да" if e["copy_operator_approved"] else "Нужно QA", labels(e["blockers"])] for e in report["entries"]],
                [12, 50, 110, 24, 20, 70])
    for row in range(2, copy_sheet.max_row + 1):
        copy_sheet.row_dimensions[row].height = 160
    extra_sheet(workbook, "Перед загрузкой",
                ["Проверка", "Действие"], [
                    ["НЕ ЗАГРУЖАТЬ ДО QA", "Файл подготовки. Ни одно объявление не допущено к публикации автоматически."],
                    ["Источник", "Read-only production snapshot " + report["captured_at"]],
                    ["Основной лист", "Только кандидаты без складских/публичных блокеров; обязательные неподтверждённые поля оставлены пустыми или выделены жёлтым."],
                    ["Полный охват", "Сверка всех устройств включает проданные, черновики и архивные; описания доступны для всех опубликованных товаров в наличии."],
                    ["Avito QA", "Подтвердить RAM, модель/память/цвет, экран/корпус, коробку и комплект. Плоский справочник не подтверждает сочетания."],
                    ["История объявлений", "Сопоставить старое объявление точно; не создавать дубль по совпадению модели или цены."],
                    ["Проданные", "Исключены. Положительный старый snapshot не отменяет статуса sold на сайте."],
                    ["IMEI и serial", "Не включены. Колонка IMEI и старые контакты/AvitoId исходного примера очищены."],
                    ["Пилот", "Текущий feed остаётся выключенным, allowlist 1-3 ID не расширялся. Полный список не разрешает массовую публикацию."],
                    ["Фотографии", "Прямые HTTPS JPEG-трансформации, без токенов, до 10 фото. WebP сайта не меняются."],
                    ["Диагностика", "Батарея и циклы из датированных августовских отчётов; перед запуском подтвердить актуальность."],
                ], [28, 115])
    workbook.active = workbook.sheetnames.index("Перед загрузкой")
    if report.get("owner_box_confirmation"):
        workbook["Перед загрузкой"].cell(6, 2).value = "Коробки открытые, Z=Нет. Подтвердить зависимые справочники и XML-комплект; заполненные RAM основаны на источниках, не на проверке каталога Avito."
    workbook.properties.description = "ISVOI Avito preparation only; not publication-ready"
    report["flat_dictionary_choice_checks"] = validate_template_choices(workbook)
    report["dependent_catalog_combinations_verified"] = False
    workbook.save(target)
    check = load_workbook(target, data_only=False)
    for ws in check:
        for row in ws:
            for cell in row:
                if cell.data_type in ("e", "f"):
                    raise ValueError("Unexpected Excel error or formula")
    for row in range(5, check[PHONE_SHEET].max_row + 1):
        if any(check[PHONE_SHEET].cell(row, col).value is not None for col in (3, 4, 5, 24, 31, 32, 33, 34)):
            raise ValueError("Historical listing/contact/private data survived")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--snapshot", required=True)
    parser.add_argument("--public-preflight", required=True)
    parser.add_argument("--approved-copy", required=True)
    parser.add_argument("--template", required=True)
    parser.add_argument("--output-dir", required=True)
    parser.add_argument("--model-reference")
    parser.add_argument("--box-confirmation")
    args = parser.parse_args()
    load = lambda name: json.loads(Path(name).read_text(encoding="utf-8"))
    snapshot, public, approved = load(args.snapshot), load(args.public_preflight), load(args.approved_copy)
    if public["captured_at"] != snapshot["captured_at"]:
        raise ValueError("Preflight belongs to another snapshot")
    report = prepare(snapshot, public, approved)
    apply_reference_evidence(report, snapshot, load(args.model_reference) if args.model_reference else None,
                             load(args.box_confirmation) if args.box_confirmation else None)
    out = Path(args.output_dir)
    out.mkdir(parents=True, exist_ok=True)
    target = out / f"ISVOI_Avito_PREPARATION_{report['captured_at'][:10]}.xlsx"
    write_workbook(Path(args.template), target, report, snapshot)
    (out / "catalog-preparation.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"audited": report["products_audited"], "published": report["published"],
                      "available_descriptions": report["available"], "catalog_candidates": report["catalog_candidate_count"],
                      "ready_to_publish": False, "workbook": str(target)}))


if __name__ == "__main__":
    main()
