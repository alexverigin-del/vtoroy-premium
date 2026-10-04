"""Prepare an explicit 1-3 device review workbook; never activates the live feed."""

import argparse
from copy import deepcopy
from datetime import datetime
import json
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill

from prepare_avito_catalog_workbook import apply_reference_evidence, extra_sheet, one, prepare, validate_template_choices, write_workbook


def normalize_product(product):
    value = deepcopy(product)
    for listing in value.get("listings") or []:
        # The current safe exporter represents an absent attrs object as {}.
        if listing.get("attributes") is None:
            listing["attributes"] = {}
    return value


def confirmed_conditions(confirmation, entries, live):
    if not confirmation:
        return {}
    if (confirmation.get("source") != "owner_chat_confirmation"
            or confirmation.get("scope") != "current_condition_battery_and_completeness_only"
            or confirmation.get("confirmed_on") != live["captured_at"][:10]):
        raise ValueError("Owner confirmation date/source/scope does not match live review")
    rows = confirmation.get("entries") or []
    if len(rows) != len(entries) or {r.get("sku") for r in rows} != {e["sku"] for e in entries}:
        raise ValueError("Owner confirmation must exactly cover selected SKU values")
    confirmed = {}
    for entry in entries:
        row = one([r for r in rows if r.get("sku") == entry["sku"]])
        product = one([p for p in live["products"] if p["sku"] == entry["sku"]])
        details = one(product["device_details"])
        if (row.get("external_id") != entry["external_id"]
                or row.get("screen_condition") != "Без дефектов"
                or row.get("case_condition") != "Без дефектов"
                or row.get("completeness") != product["completeness"]
                or row.get("battery") != details["battery"]
                or row.get("battery_cycles") != details["battery_cycles"]):
            raise ValueError("Owner confirmation does not match device facts: " + entry["sku"])
        confirmed[entry["sku"]] = row
    return confirmed


def select_pilot(snapshot, live, public, approved, skus, confirmation=None):
    if not 1 <= len(skus) <= 3 or len(set(skus)) != len(skus):
        raise ValueError("Select 1-3 distinct SKU values")
    if public.get("captured_at") != snapshot.get("captured_at"):
        raise ValueError("Public evidence belongs to another snapshot")
    if (not live.get("captured_at") or datetime.fromisoformat(live["captured_at"])
            < datetime.fromisoformat(snapshot["captured_at"])):
        raise ValueError("Live evidence is missing or older than source snapshot")
    if live.get("private_identifier_redactions"):
        raise ValueError("Redacted live source requires operator review")
    for sku in skus:
        previous = one([p for p in snapshot["products"] if p["sku"] == sku])
        current = one([p for p in live["products"] if p["sku"] == sku])
        if not previous or not current:
            raise ValueError("Missing or ambiguous pilot SKU: " + sku)
        if normalize_product(previous) != normalize_product(current):
            raise ValueError("Pilot facts changed; refresh public evidence and copy: " + sku)
    previous_store = one([s for s in snapshot["stores"] if s["slug"] == "belgorod"])
    current_store = one([s for s in live["stores"] if s["slug"] == "belgorod"])
    if not previous_store or previous_store != current_store:
        raise ValueError("Store facts changed; refresh preparation")

    # Assess the full current catalog before narrowing it, so out-of-pilot
    # collisions are still blockers. Public HTTP evidence is explicitly reused.
    report = prepare(live, public, approved)
    selected = []
    for sku in skus:
        entry = one([e for e in report["entries"] if e["sku"] == sku])
        if not entry or not entry["catalog_checks_pass"]:
            raise ValueError("SKU is not a catalog candidate: " + sku)
        selected.append(entry)
    conditions = confirmed_conditions(confirmation, selected, live)
    for entry in selected:
        if entry["sku"] in conditions:
            entry["confirmed_phone_attributes"] = {"ScreenCondition": "Без дефектов", "CaseCondition": "Без дефектов"}
            entry["warnings"] = [w for w in entry["warnings"] if w != "diagnostic_is_historical_recheck_before_publication"]
            for audit in report["audits"]:
                if audit["sku"] == entry["sku"]:
                    audit["warnings"] = entry["warnings"]
    report.update(
        entries=selected,
        whole_catalog_available=report["available"],
        whole_catalog_candidate_count=report["catalog_candidate_count"],
        available=len(selected),
        catalog_candidate_count=len(selected),
        selection_scope="pilot_preparation_only",
        selected_skus=list(skus),
        proposed_pilot_external_ids=[e["external_id"] for e in selected],
        live_allowlist_changed=False,
        public_evidence_captured_at=public["captured_at"],
        public_http_checks_repeated=False,
        unchanged_pilot_facts_verified=True,
        historical_avito_listing_reconciled=False,
        current_state_owner_confirmed=bool(conditions),
        physical_recheck_performed_by_agent=False,
        owner_confirmation=confirmation,
        official_catalog_combinations_confirmed=False,
        feed_activated=False,
    )
    return report


def write_pilot(template, target, report, snapshot):
    write_workbook(template, target, report, snapshot)
    workbook = load_workbook(target)
    workbook["Все доступные описания"].title = "Описания пилота"
    for row in range(2, workbook["Описания пилота"].max_row + 1):
        workbook["Описания пилота"].row_dimensions[row].height = 330
    phone = workbook["Телефоны-Мобильные телефоны"]
    for index, entry in enumerate(report["entries"], 5):
        for column, key in ((29, "ScreenCondition"), (30, "CaseCondition")):
            value = entry.get("confirmed_phone_attributes", {}).get(key)
            if value:
                phone.cell(index, column).value = value
                phone.cell(index, column).fill = PatternFill("solid", fgColor="E2F0D9")
                phone.cell(index, column).comment = Comment(
                    "Источник: подтверждение владельца в чате, " + report["owner_confirmation"]["confirmed_on"]
                    + "; SKU " + entry["sku"] + ". Не подтверждает XML-схему или публикацию.", "ISVOI")
    before = workbook["Перед загрузкой"]
    before.cell(5, 2).value = (
        "Основной лист и описания ограничены выбранными SKU. Полный аудит сохранён отдельно; "
        "он не расширяет пилот. Ни одна строка не готова к публикации автоматически."
    )
    before.cell(10, 2).value = (
        "Пилот подготовки: " + ", ".join(report["selected_skus"])
        + ". Live allowlist не изменён, feed выключен. Это не разрешение на загрузку."
    )
    if report["current_state_owner_confirmed"]:
        before.cell(6, 2).value = "Состояние экрана/корпуса и комплект подтверждены владельцем. Подтвердить справочные сочетания, RAM и условные XML-поля."
        before.cell(12, 2).value = "Владелец подтвердил неизменность батареи; даты исходных августовских отчётов сохранены, новый сертификат не создавался."
    if report.get("owner_box_confirmation"):
        before.cell(6, 2).value = "Экран/корпус и комплект подтверждены. Коробки открытые, Z=Нет. RAM заполнена по источникам. Остались справочные сочетания и XML/кабинетный QA."
    rows = []
    for entry in report["entries"]:
        product = next(p for p in snapshot["products"] if p["id"] == entry["product_id"])
        details = one(product["device_details"])
        confirmed = all(key in entry.get("confirmed_phone_attributes", {})
                        for key in ("ScreenCondition", "CaseCondition"))
        ram = entry.get("model_reference_attributes", {}).get("RamSize")
        opened_box = entry.get("confirmed_phone_attributes", {}).get("BoxSealed") == "Нет"
        rows.extend([
            [entry["sku"], "Модель / память / цвет", product["model"] + " / " + details["storage"]
             + " / " + product["color"], "Проверить в зависимом справочнике Avito; плоский список недостаточен"],
            [entry["sku"], "Оперативная память (RAM)", ram or "Не заполнено",
             "Источник: " + entry["model_reference_source"] + ". Подтвердить сочетание в Avito"
             if ram else "Выбрать точное значение для модели в Avito; не путать с накопителем"],
            [entry["sku"], "Экран / корпус", "Без дефектов / без дефектов" if confirmed else "Грейд " + details["grade"],
             "Подтверждено владельцем " + report["owner_confirmation"]["confirmed_on"] + "; AC / AD заполнены"
             if confirmed else "Подтвердить актуальное состояние и заполнить AC / AD. Проверка экрана не доказывает отсутствие царапин"],
            [entry["sku"], "Комплект / коробка", product["completeness"],
             "Владелец подтвердил открытую коробку, Z=Нет. XML-комплект ещё требует проверки; блок питания не добавлялся"
             if opened_box else "Коробка и кабель подтверждены; запечатанность не заявлена. Проверить условное поле Z, зарядный блок не добавлять"],
            [entry["sku"], "Батарея", str(details["battery"]) + "; " + str(details["battery_cycles"])
             + " циклов; отчёт " + details["diagnostic_date"], "Владелец подтвердил, что значения не изменились. Дата отчёта сохранена"
             if confirmed else "Августовский отчёт, не новое измерение. Подтвердить актуальность"],
            [entry["sku"], "Сбалансированное описание", "Ранее одобрено" if entry["copy_operator_approved"] else "Нужно одобрение текста",
             "Одобрение только текста, не публикации; для нового SKU не наследовать одобрение другого устройства"],
            [entry["sku"], "Цена / остаток", str(entry["price_rub"]) + " ₽ / " + str(entry["stock_quantity"]) + " шт.",
             "Read-only сверка " + report["captured_at"] + ". Повторить непосредственно перед запуском"],
            [entry["sku"], "Стабильный Id", entry["external_id"], "Существующий ID канала сохранён, новый не создавался"],
        ])
    rows.extend([
        ["Все", "История Avito", "Не сопоставлена", "Точно связать старое объявление со SKU либо документировать отсутствие пересечения. Модель и цена не являются ключом"],
        ["Все", "XML и кабинет", "Не проверены", "Подтвердить зависимости и XML; проверить источники и расписание кабинета. Не запускать загрузку автоматически"],
        ["Все", "Публичные страницы и фото", "Повторно HTTP не запрашивались", "Успешные проверки от " + report["public_evidence_captured_at"]
         + " переиспользованы после проверки неизменности всех фактов выбранных устройств"],
    ])
    qa = extra_sheet(workbook, "Проверки пилота", ["SKU", "Поле", "Что известно", "Что осталось"], rows, [14, 30, 68, 110])
    for row in range(2, qa.max_row + 1):
        qa.row_dimensions[row].height = 48
    workbook.active = workbook.sheetnames.index("Проверки пилота")
    workbook.properties.description = "ISVOI 1-3 device pilot preparation only; not publication-ready"
    report["flat_dictionary_choice_checks"] = validate_template_choices(workbook)
    workbook.save(target)
    check = load_workbook(target)
    phone = check["Телефоны-Мобильные телефоны"]
    ids = [phone.cell(row, 1).value for row in range(5, phone.max_row + 1)]
    if ids != report["proposed_pilot_external_ids"]:
        raise ValueError("Pilot workbook ID set/order differs from explicit selection")
    if len(phone.data_validations.dataValidation) != 18:
        raise ValueError("Official template dropdowns were not preserved")
    for sheet in check:
        for row in sheet:
            if any(cell.data_type in ("f", "e") for cell in row):
                raise ValueError("Unexpected formula or Excel error")


def main():
    parser = argparse.ArgumentParser()
    for name in ("snapshot", "live-snapshot", "public-preflight", "approved-copy", "template", "output-dir"):
        parser.add_argument("--" + name, required=True)
    parser.add_argument("--sku", action="append", required=True)
    parser.add_argument("--operator-confirmation")
    parser.add_argument("--model-reference")
    parser.add_argument("--box-confirmation")
    args = parser.parse_args()
    load = lambda path: json.loads(Path(path).read_text(encoding="utf-8"))
    snapshot, live = load(args.snapshot), load(args.live_snapshot)
    confirmation = load(args.operator_confirmation) if args.operator_confirmation else None
    report = select_pilot(snapshot, live, load(args.public_preflight), load(args.approved_copy), args.sku, confirmation)
    apply_reference_evidence(report, live, load(args.model_reference) if args.model_reference else None,
                             load(args.box_confirmation) if args.box_confirmation else None)
    out = Path(args.output_dir)
    out.mkdir(parents=True, exist_ok=True)
    target = out / f"ISVOI_Avito_PILOT_PREPARATION_{report['captured_at'][:10]}.xlsx"
    write_pilot(Path(args.template), target, report, live)
    (out / "pilot-review.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"selected_skus": report["selected_skus"], "rows": len(report["entries"]),
                      "ready_to_publish": False, "workbook": str(target)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
